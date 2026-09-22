"""U-021 steps 2-3: bring the Square item library up to date with Sellorama.

Square was filled from Sellorama in June and has not been used since, so its prices are
three months old, and it cannot scan Aaojee labels (6-digit SKU, no GTIN; the scanner
sends the full 8-digit UPC-E).  The owner has also edited Square by hand since June
(names, categories, tax, a few items added only in Square), so this script changes as
little as possible:

  * GTIN   - every item whose SKU is a 6-digit code and whose GTIN is blank gets the
             8-digit scan code (314300 -> 03143009).  Confirmed working by TEST_A.
  * Price  - taken from Sellorama when it differs (Sellorama is where prices are kept),
             unless the price was changed by hand in Square since the June import.
  * New    - Sellorama items whose barcode is not in Square are added as new items.
  Names, categories, tax and everything else in Square are left alone.  Differences
  there are listed in the review workbook instead.

Inputs (read only, never modified):
  square_export_before_U-021.xlsx        Square export (Items -> Actions -> Export library)
  ../../store_sync/Products.accdb        copy of Sellorama's database (program closed)
  ../../september_store_copy/READY_FOR_STORE/products.db   label program database
  ../../POS Upgrade Files for Context/.../square_catalog_IMPORT.xlsx   the June import file

Outputs (in this folder):
  square_UPDATE_U-021.xlsx   the file to import into Square (update, never replace)
  changes.csv                every change the import makes
  REVIEW_U-021.xlsx          things for the owner to decide; nothing in it is applied

    python build_square_update.py                  # build the files above
    python build_square_update.py --check NEW.xlsx # after importing: compare a new export

Needs:  pip install openpyxl access-parser
"""
import argparse
import collections
import csv
import difflib
import re
import sqlite3
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT / "source"))
from barcode_engine import compute_check_digit  # noqa: E402

EXPORT = HERE / "square_export_before_U-021.xlsx"
SELLORAMA = ROOT / "store_sync" / "Products.accdb"
LABEL_DB = ROOT / "september_store_copy" / "READY_FOR_STORE" / "products.db"
JUNE_IMPORT = (ROOT / "POS Upgrade Files for Context" / "SelloramaPOS Product Info Transfer"
               / "Square Export" / "square_catalog_IMPORT.xlsx")
OUT_UPDATE = HERE / "square_UPDATE_U-021.xlsx"
OUT_CHANGES = HERE / "changes.csv"
OUT_REVIEW = HERE / "REVIEW_U-021.xlsx"

HEADER_ROW = 2
TAX = "Tax - Sales Tax (8.125%)"
# Sellorama rows that must never reach Square
EXCLUDE_BARCODES = {
    "012800517725": "RAYOVAC CLAUDE TEST - a test item, not a product",
}
# 6-digit SKUs that are placeholders, not label codes
NO_SCAN_CODE = {"000000"}
# Settings for new items: the values nearly every existing Square item has
NEW_ITEM_DEFAULTS = {
    "Square Online Item Visibility": "unavailable",
    "Item Type": "Physical good",
    "Pickup Enabled": "Y",
    "Archived": "N",
    "Contains Alcohol": "N",
    TAX: "N",
}


def s(v) -> str:
    return "" if v is None else re.sub(r"\s+", " ", str(v)).strip()


def money(v):
    try:
        return round(float(str(v).replace("$", "").replace(",", "")), 2)
    except (TypeError, ValueError):
        return None


def scan_code(data6: str) -> str:
    return "0" + data6 + str(compute_check_digit(data6))


def gs1_ok(code: str) -> bool:
    body, check = code[:-1], int(code[-1])
    total = sum(int(c) * (3 if i % 2 == 0 else 1) for i, c in enumerate(reversed(body)))
    return (10 - total % 10) % 10 == check


def gtin_for(barcode: str) -> str:
    """GTIN for a Sellorama barcode, or '' if it has none Square can use."""
    if not barcode.isdigit():
        return ""
    if len(barcode) == 6:
        return scan_code(barcode)
    if len(barcode) == 8 and (gs1_ok(barcode) or
                              (barcode[0] == "0" and int(barcode[7]) == compute_check_digit(barcode[1:7]))):
        return barcode
    if len(barcode) in (12, 13, 14) and gs1_ok(barcode):
        return barcode
    return ""


# ── Loading ──────────────────────────────────────────────────────────────────

def load_export(path: Path):
    wb = openpyxl.load_workbook(path, read_only=True)
    rows = list(wb["Items"].iter_rows(values_only=True))
    headers = list(rows[HEADER_ROW - 1])
    data = [dict(zip(headers, r)) for r in rows[HEADER_ROW:] if any(v not in (None, "") for v in r)]
    return headers, data


def load_june_import():
    """SKU -> row of the file Square was filled from in June (headers on row 6)."""
    wb = openpyxl.load_workbook(JUNE_IMPORT, read_only=True)
    rows = list(wb["Items"].iter_rows(values_only=True))
    headers = list(rows[5])
    return {s(r[headers.index("SKU")]): dict(zip(headers, r)) for r in rows[6:]
            if any(v not in (None, "") for v in r) and s(r[headers.index("SKU")])}


def load_sellorama():
    from access_parser import AccessParser
    table = AccessParser(str(SELLORAMA)).parse_table("barcode")
    cols = list(table)
    return [{c: table[c][i] for c in cols} for i in range(len(table[cols[0]]))]


def load_labels():
    db = sqlite3.connect(f"file:{LABEL_DB.as_posix()}?mode=ro&immutable=1", uri=True)
    return db.execute("SELECT name, barcode_number, price FROM products "
                      "WHERE deleted_at IS NULL AND barcode_number <> '' ORDER BY name").fetchall()


# ── Building ─────────────────────────────────────────────────────────────────

def build():
    headers, sq = load_export(EXPORT)
    sello = load_sellorama()
    labels = load_labels()
    june = load_june_import()

    sq_by_sku = {s(d["SKU"]): d for d in sq if s(d["SKU"])}
    sel_by_bc = {}
    for r in sello:
        b = s(r["Barcode"])
        if b:
            if b in sel_by_bc:
                sys.exit(f"Sellorama has barcode {b} twice; sort that out first")
            sel_by_bc[b] = r
    categories = {s(d["Categories"]).upper(): s(d["Categories"]) for d in sq if s(d["Categories"])}
    sq_names = {s(d["Item Name"]).upper() for d in sq}

    changes = []           # (token, item, variation, sku, field, old, new, reason)
    edits = {}             # token -> {field: new value}
    review = collections.defaultdict(list)

    def change(d, field, new, reason):
        edits.setdefault(d["Token"], {})[field] = new
        changes.append((d["Token"], s(d["Item Name"]), s(d["Variation Name"]), s(d["SKU"]),
                        field, s(d[field]), new, reason))

    # 1. Scan codes for 6-digit SKUs
    for d in sq:
        sku = s(d["SKU"])
        if len(sku) == 6 and sku.isdigit() and not s(d["GTIN"]):
            if sku in NO_SCAN_CODE:
                review["Scan code not added"].append((sku, s(d["Item Name"]), "placeholder code"))
                continue
            change(d, "GTIN", scan_code(sku), "scanner sends the 8-digit code")

    # 2. Prices, and differences that are only reported
    for b, r in sel_by_bc.items():
        d = sq_by_sku.get(b)
        if not d:
            continue
        new_p, old_p = money(r["SellingRate"]), money(d["Price"])
        if new_p != old_p:
            j = june.get(b)
            if not new_p:
                review["Price not applied"].append((b, s(d["Item Name"]), old_p, s(r["SellingRate"]),
                                                   "Sellorama price is zero or blank"))
            elif j is None or money(j["Price"]) != old_p:
                review["Price not applied"].append((b, s(d["Item Name"]), old_p, new_p,
                                                   "price was changed in Square after the June import"))
            else:
                change(d, "Price", new_p, "Sellorama price")
        if s(r["ItemName"]).upper() != s(d["Item Name"]).upper():
            review["Names differ"].append((b, s(d["Item Name"]), s(r["ItemName"])))

    # 3. New items
    new_rows = []
    for b, r in sel_by_bc.items():
        if b in sq_by_sku:
            continue
        name, price = s(r["ItemName"]), money(r["SellingRate"])
        if b in EXCLUDE_BARCODES:
            review["New items not added"].append((b, name, price, EXCLUDE_BARCODES[b]))
            continue
        if not price:
            review["New items not added"].append((b, name, s(r["SellingRate"]), "price is zero or blank"))
            continue
        dept = s(r["Department"])
        cat = categories.get(dept.upper(), dept)
        gtin = gtin_for(b)
        row = {h: "" for h in headers}
        row.update(NEW_ITEM_DEFAULTS)
        row.update({"Item Name": name, "Variation Name": s(r["ItemSize"]) or "Regular", "SKU": b,
                    "GTIN": gtin, "Price": price, "Categories": cat, "Reporting Category": cat})
        new_rows.append(row)
        notes = []
        if not gtin:
            notes.append("no valid GTIN: scans only if Square matches the SKU")
        if name.upper() in sq_names:
            notes.append("an item with this name is already in Square")
        review["New items"].append((b, name, row["Variation Name"], price, cat, gtin, "; ".join(notes)))
        changes.append(("(new)", name, row["Variation Name"], b, "new item", "", f"{price} / {cat}",
                        "in Sellorama, not in Square"))

    # 4. Items only in Square (not in Sellorama)
    blank_names = {s(r["ItemName"]).upper() for r in sello if not s(r["Barcode"])}
    no_barcode = 0
    for d in sq:
        sku = s(d["SKU"])
        if sku and sku not in sel_by_bc:
            if s(d["Item Name"]).upper() in blank_names:
                no_barcode += 1      # Sellorama has it with no barcode; Square made up a SKU
                continue
            review["Only in Square"].append((sku, s(d["Item Name"]), s(d["Variation Name"]), money(d["Price"]),
                                             "not in Sellorama (added in Square?)"))

    # 5. Label program against Sellorama (what a printed label rings up)
    norm = lambda x: re.sub("[^A-Z]", "", x.upper())
    for name, code, price in labels:
        r = sel_by_bc.get(code)
        if not r:
            review["Labels not in Sellorama"].append((code, name, price, scan_code(code)))
            continue
        sname = s(r["ItemName"])
        a, b2 = norm(name), norm(sname)
        if difflib.SequenceMatcher(None, a, b2).ratio() < 0.75 and a not in b2 and b2 not in a:
            review["Label code = other product"].append((code, name, price, sname, money(r["SellingRate"])))
        if money(price) != money(r["SellingRate"]):
            review["Label price differs"].append((code, name, money(price), money(r["SellingRate"])))

    # Collision check: no scan code or new barcode may already belong to another item
    taken = collections.Counter()
    for d in sq:
        for v in {s(d["SKU"]), s(d["GTIN"])} - {""}:
            taken[v] += 1
    for tok, e in edits.items():
        if "GTIN" in e and taken[e["GTIN"]]:
            sys.exit(f"Scan code {e['GTIN']} already used in Square")
        if "GTIN" in e:
            taken[e["GTIN"]] += 1
    for row in new_rows:
        for v in {row["SKU"], row["GTIN"]} - {""}:
            if taken[v]:
                sys.exit(f"New item code {v} ({row['Item Name']}) already used in Square")
            taken[v] += 1

    write_update(headers, sq, edits, new_rows)
    write_changes(changes)
    write_review(review)

    kinds = collections.Counter(c[4] for c in changes)
    print(f"Square rows: {len(sq)}   Sellorama rows: {len(sello)}   label products: {len(labels)}")
    print(f"Changes: {dict(kinds)}   items touched: {len(edits)}")
    print(f"  ({no_barcode} Square items are Sellorama items with no barcode: nothing to do)")
    for sheet, items in review.items():
        print(f"  review - {sheet}: {len(items)}")
    print(f"Wrote {OUT_UPDATE.name}, {OUT_CHANGES.name}, {OUT_REVIEW.name}")


def write_update(headers, sq, edits, new_rows):
    """All rows of every touched item (Square keeps item-level settings on the first row
    of an item, so an item is never sent half), edited cells changed, plus new items."""
    touched_names = {s(d["Item Name"]) for d in sq if d["Token"] in edits}
    out = []
    for d in sq:
        if s(d["Item Name"]) in touched_names:
            row = dict(d)
            row.update(edits.get(d["Token"], {}))
            out.append(row)
    out.extend(new_rows)

    wb = openpyxl.load_workbook(EXPORT)
    ws = wb["Items"]
    ws.delete_rows(HEADER_ROW + 1, ws.max_row - HEADER_ROW)
    text_cols = {headers.index("SKU") + 1, headers.index("GTIN") + 1}
    for i, row in enumerate(out, start=HEADER_ROW + 1):
        for c, h in enumerate(headers, start=1):
            v = row.get(h)
            cell = ws.cell(i, c, "" if v is None else v)
            if c in text_cols:
                cell.value = s(v)
                cell.number_format = "@"
    wb.save(OUT_UPDATE)


def write_changes(changes):
    with open(OUT_CHANGES, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["token", "item", "variation", "sku", "field", "old", "new", "reason"])
        w.writerows(changes)


REVIEW_HEADERS = {
    "Label code = other product": ["code", "label program name", "label price", "Sellorama name", "Sellorama price"],
    "Label price differs": ["code", "product", "label price", "Sellorama price"],
    "Labels not in Sellorama": ["code", "product", "label price", "scan code"],
    "New items": ["barcode", "name", "size", "price", "category", "GTIN", "note"],
    "New items not added": ["barcode", "name", "price", "why"],
    "Names differ": ["barcode", "name in Square (kept)", "name in Sellorama"],
    "Only in Square": ["SKU", "item", "variation", "price", "why"],
    "Price not applied": ["barcode", "item", "Square price", "Sellorama price", "why"],
    "Scan code not added": ["SKU", "item", "why"],
}


def write_review(review):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, head in REVIEW_HEADERS.items():
        if not review.get(sheet):
            continue
        ws = wb.create_sheet(sheet[:31])
        ws.append(head + ["your decision"])
        for row in review[sheet]:
            ws.append(list(row))
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(45, max(10, *(len(s(c.value)) + 2 for c in col)))
        ws.freeze_panes = "A2"
    wb.save(OUT_REVIEW)


# ── Checking an export taken after the import ────────────────────────────────

def check(new_export: Path):
    headers, before = load_export(EXPORT)
    _, after = load_export(new_export)
    expected = collections.defaultdict(dict)
    new_expected = {}
    with open(OUT_CHANGES, encoding="utf-8-sig") as f:
        for c in csv.DictReader(f):
            if c["token"] == "(new)":
                new_expected[c["sku"]] = c
            else:
                expected[c["token"]][c["field"]] = c["new"]
    before_by_tok = {d["Token"]: d for d in before}
    after_by_tok = {d["Token"]: d for d in after}
    problems = []
    for tok, b in before_by_tok.items():
        a = after_by_tok.get(tok)
        if a is None:
            problems.append(f"missing after import: {s(b['Item Name'])} ({s(b['SKU'])})")
            continue
        for h in headers:
            want = expected[tok].get(h, s(b[h]))
            got = s(a[h])
            if h == "Price":
                same = money(want) == money(got)
            else:
                same = want == got
            if not same:
                problems.append(f"{s(b['Item Name'])} ({s(b['SKU'])}) {h}: expected {want!r}, found {got!r}")
    added = [d for tok, d in after_by_tok.items() if tok not in before_by_tok]
    added_skus = {s(d["SKU"]) for d in added}
    for sku, c in new_expected.items():
        if sku not in added_skus:
            problems.append(f"new item not found: {c['item']} ({sku})")
    for d in added:
        if s(d["SKU"]) not in new_expected:
            problems.append(f"unexpected new item: {s(d['Item Name'])} ({s(d['SKU'])})")
    print(f"Checked {len(before_by_tok)} existing rows and {len(added)} new rows.")
    print("\n".join(problems[:200]) if problems else "Everything matches: only the planned changes were made.")
    if len(problems) > 200:
        print(f"... and {len(problems) - 200} more")
    return not problems


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--check", type=Path, help="Square export taken after the import")
    args = ap.parse_args()
    if args.check:
        sys.exit(0 if check(args.check) else 1)
    build()
