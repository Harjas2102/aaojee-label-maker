"""U-021 step 1: one-item Square import files to find which field makes a label scan.

Square stores in-house items under their 6-digit code (SKU 314300, GTIN blank), but the
scanner sends the full 8-digit UPC-E (03143009), so Square finds nothing.  This builds two
one-row copies of the Square library export for ALOO MATTER:

  TEST_A_gtin.xlsx  GTIN = 03143009, SKU stays 314300   (try this first)
  TEST_B_sku.xlsx   SKU  = 03143009, GTIN stays blank   (only if Square rejects A)

Every other cell of the row, including Token, is copied unchanged, so Square updates the
existing item instead of creating a new one.

    python make_test_files.py
"""
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[1] / "source"))
from barcode_engine import compute_check_digit  # noqa: E402

EXPORT = HERE / "square_export_before_U-021.xlsx"
HEADER_ROW = 2          # Square exports: row 1 blank, row 2 headers, data from row 3
TEST_SKU = "314300"     # ALOO MATTER


def scan_code(data6: str) -> str:
    return "0" + data6 + str(compute_check_digit(data6))


def build(field: str, out_name: str) -> None:
    wb = openpyxl.load_workbook(EXPORT)
    ws = wb["Items"]
    headers = [c.value for c in ws[HEADER_ROW]]
    col = {h: i + 1 for i, h in enumerate(headers)}

    matches = [r for r in range(HEADER_ROW + 1, ws.max_row + 1)
               if str(ws.cell(r, col["SKU"]).value or "").strip() == TEST_SKU]
    if len(matches) != 1:
        sys.exit(f"Expected exactly one row with SKU {TEST_SKU}, found {len(matches)}")
    src = matches[0]
    first = HEADER_ROW + 1

    if src != first:
        for c in range(1, ws.max_column + 1):
            ws.cell(first, c).value = ws.cell(src, c).value
    ws.delete_rows(first + 1, ws.max_row - first)

    code = scan_code(TEST_SKU)
    ws.cell(first, col[field]).value = code
    ws.cell(first, col[field]).number_format = "@"
    wb.save(HERE / out_name)

    # Read back and show exactly what the file contains
    check = openpyxl.load_workbook(HERE / out_name, read_only=True)["Items"]
    rows = [r for r in check.iter_rows(min_row=first, values_only=True) if any(v not in (None, "") for v in r)]
    assert len(rows) == 1, rows
    row = dict(zip(headers, rows[0]))
    print(f"{out_name}: {row['Item Name']} | Token {row['Token']} | SKU {row['SKU']!r} | "
          f"GTIN {row['GTIN']!r} | Price {row['Price']} | Tax {row['Tax - Sales Tax (8.125%)']}")


if __name__ == "__main__":
    build("GTIN", "TEST_A_gtin.xlsx")
    build("SKU", "TEST_B_sku.xlsx")
